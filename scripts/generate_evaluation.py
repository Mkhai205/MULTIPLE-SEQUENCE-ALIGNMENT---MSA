#!/usr/bin/env python3
"""
Team Member Evaluation Report Generator for Group 5 MSA Project
Generates: Ban_Danh_Gia_Thanh_Vien_Nhom_5.docx
Standard 10-Point Grading Scale, Work Assignment, Cross-Evaluation Matrix, and Signatures.
"""

import os
from pathlib import Path
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_ALIGN_VERTICAL
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT_DIR = Path(__file__).resolve().parent.parent
OUTPUT_DOCX = ROOT_DIR / "Ban_Danh_Gia_Thanh_Vien_Nhom_5.docx"

def set_cell_background(cell, hex_color):
    """Sets background color of a table cell."""
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement('w:shd')
    shd.set(qn('w:val'), 'clear')
    shd.set(qn('w:color'), 'auto')
    shd.set(qn('w:fill'), hex_color)
    tc_pr.append(shd)

def set_cell_margins(cell, top=100, bottom=100, left=150, right=150):
    """Sets cell margins in twips."""
    tc_pr = cell._tc.get_or_add_tcPr()
    tc_mar = OxmlElement('w:tcMar')
    for m, val in [('w:top', top), ('w:bottom', bottom), ('w:left', left), ('w:right', right)]:
        node = OxmlElement(m)
        node.set(qn('w:w'), str(val))
        node.set(qn('w:type'), 'dxa')
        tc_mar.append(node)
    tc_pr.append(tc_mar)

def create_evaluation_doc():
    doc = Document()

    # Set page margins
    sections = doc.sections
    for section in sections:
        section.top_margin = Inches(0.8)
        section.bottom_margin = Inches(0.8)
        section.left_margin = Inches(0.9)
        section.right_margin = Inches(0.9)

    # Styles
    normal_style = doc.styles['Normal']
    normal_style.font.name = 'Times New Roman'
    normal_style.font.size = Pt(12)
    normal_style.font.color.rgb = RGBColor(15, 23, 42)

    # Header / University
    p_uni = doc.add_paragraph()
    p_uni.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r1 = p_uni.add_run("BỘ GIÁO DỤC VÀ ĐÀO TẠO • TRƯỜNG ĐẠI HỌC CÔNG NGHỆ\n")
    r1.font.bold = True
    r1.font.size = Pt(12)
    r2 = p_uni.add_run("KHOA CÔNG NGHỆ THÔNG TIN - BỘ MÔN KHOA HỌC MÁY TÍNH\n")
    r2.font.bold = True
    r2.font.size = Pt(12)
    r3 = p_uni.add_run("------------------------------------------------------------------")
    r3.font.color.rgb = RGBColor(100, 116, 139)

    # Document Title
    p_title = doc.add_paragraph()
    p_title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    p_title.paragraph_format.space_before = Pt(14)
    p_title.paragraph_format.space_after = Pt(4)
    run_t = p_title.add_run("BẢN ĐÁNH GIÁ THÀNH VIÊN VÀ MỨC ĐỘ ĐÓNG GÓP\n")
    run_t.font.bold = True
    run_t.font.size = Pt(16)
    run_t.font.color.rgb = RGBColor(16, 44, 87)

    run_sub = p_title.add_run("(Áp dụng Thang điểm 10 theo Quy định Đánh giá Bài tập lớn Học phần)\n")
    run_sub.font.italic = True
    run_sub.font.size = Pt(11)
    run_sub.font.color.rgb = RGBColor(71, 85, 105)

    # Subject & Topic Metadata
    p_meta = doc.add_paragraph()
    p_meta.paragraph_format.space_after = Pt(14)
    p_meta.paragraph_format.line_spacing = 1.25

    meta_items = [
        ("Học phần:", " Phân tích và Thiết kế Thuật toán"),
        ("Tên Đề tài:", " Thuật toán Chia để trị Song song kết hợp Quy hoạch động cho Bài toán Căn chỉnh Đa chuỗi Gen (Multiple Sequence Alignment - MSA)"),
        ("Nhóm Thực hiện:", " NHÓM 5 (Gồm 05 Sinh viên)"),
        ("Thời gian Thực hiện:", " Học kỳ I - Năm học 2025 - 2026")
    ]
    for lbl, val in meta_items:
        r = p_meta.add_run(lbl)
        r.font.bold = True
        p_meta.add_run(val + "\n")

    # -------------------------------------------------------------------------
    # PHẦN 1: TIÊU CHÍ ĐÁNH GIÁ (THANG ĐIỂM 10)
    # -------------------------------------------------------------------------
    h1 = doc.add_paragraph()
    h1.paragraph_format.space_before = Pt(12)
    h1.paragraph_format.space_after = Pt(6)
    r = h1.add_run("I. KHUNG TIÊU CHÍ VÀ THANG ĐIỂM ĐÁNH GIÁ (THANG ĐIỂM 10)")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(16, 44, 87)

    p_crit_intro = doc.add_paragraph(
        "Nhóm 5 thống nhất xây dựng khung tiêu chí đánh giá minh bạch, công bằng dựa trên 5 khía cạnh cốt lõi với thang điểm 10 chuẩn mực như sau:"
    )
    p_crit_intro.paragraph_format.space_after = Pt(8)
    p_crit_intro.paragraph_format.line_spacing = 1.25

    # Table 1: Criteria table
    t1 = doc.add_table(rows=6, cols=4)
    t1.alignment = WD_TABLE_ALIGNMENT.CENTER
    t1.autofit = False

    t1_headers = ["STT", "Tiêu Chí Đánh Giá", "Trọng Số (%)", "Điểm Tối Đa"]
    t1_widths = [Inches(0.6), Inches(4.3), Inches(1.2), Inches(1.1)]

    for c_idx, head in enumerate(t1_headers):
        cell = t1.cell(0, c_idx)
        cell.width = t1_widths[c_idx]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(head)
        r.font.bold = True
        r.font.size = Pt(10.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    crit_rows = [
        ("1", "Ý thức Chuyên cần & Tinh thần Trách nhiệm:\nTham gia đầy đủ các buổi họp nhóm, đóng góp thảo luận tích cực, tuân thủ kỷ luật chung.", "20%", "2.0"),
        ("2", "Khối lượng Công việc Đảm nhiệm & Tiến độ:\nHoàn thành 100% nhiệm vụ được giao đúng deadline, chủ động nhận thêm các task phát sinh.", "30%", "3.0"),
        ("3", "Chất lượng Kỹ thuật & Độ Chính xác:\nMã nguồn chạy ổn định, thuật toán chuẩn xác, tối ưu hóa tốt, vượt qua toàn bộ các bộ test case.", "25%", "2.5"),
        ("4", "Kỹ năng Phối hợp & Làm việc Nhóm:\nGiao tiếp hiệu quả, hỗ trợ thành viên khác xử lý lỗi, sử dụng Git thành thạo, không gây xung đột mã nguồn.", "15%", "1.5"),
        ("5", "Đóng góp Sáng tạo & Xử lý Vấn đề:\nĐưa ra giải pháp cải tiến thuật toán, giải quyết triệt để các rào cản kỹ thuật phức tạp.", "10%", "1.0")
    ]

    for r_idx, row_data in enumerate(crit_rows, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = t1.cell(r_idx, c_idx)
            cell.width = t1_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 100, 100)
            p = cell.paragraphs[0]
            if c_idx in [0, 2, 3]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(10)
            if c_idx == 3:
                r.font.bold = True

    # -------------------------------------------------------------------------
    # PHẦN 2: PHÂN CÔNG NHIỆM VỤ CHI TIẾT
    # -------------------------------------------------------------------------
    h2 = doc.add_paragraph()
    h2.paragraph_format.space_before = Pt(16)
    h2.paragraph_format.space_after = Pt(6)
    r = h2.add_run("II. PHÂN CÔNG NHIỆM VỤ VÀ KẾT QUẢ ĐẠT ĐƯỢC CỦA 5 THÀNH VIÊN")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(16, 44, 87)

    t2 = doc.add_table(rows=6, cols=5)
    t2.alignment = WD_TABLE_ALIGNMENT.CENTER
    t2.autofit = False

    t2_headers = ["STT", "Họ và Tên", "MSSV", "Vai Trò & Nhiệm Vụ Phân Công Đảm Nhiệm", "Kết Quả Bàn Giao"]
    t2_widths = [Inches(0.5), Inches(1.5), Inches(1.0), Inches(2.8), Inches(1.4)]

    for c_idx, head in enumerate(t2_headers):
        cell = t2.cell(0, c_idx)
        cell.width = t2_widths[c_idx]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(head)
        r.font.bold = True
        r.font.size = Pt(10)
        r.font.color.rgb = RGBColor(255, 255, 255)

    members_tasks = [
        ("1", "Nguyễn Văn An", "20210001",
         "Trưởng nhóm • Thiết kế kiến trúc tổng thể pipeline C++17 modular; Nghiên cứu lý thuyết Gotoh và cài đặt thuật toán Hirschberg / Myers-Miller O(N); Xây dựng Profile-Profile Aligner.",
         "Mã nguồn core/ & align/\n100% Score match\nHoàn thành Xuất sắc"),
        ("2", "Trần Thị Bình", "20210002",
         "Thành viên • Thiết kế kiến trúc song song OpenMP; Phẳng hóa vòng lặp All-Pairs và cấu hình Dynamic Scheduling; Cài đặt thuật toán lập lịch Wavefront Tree Scheduler trên cây UPGMA.",
         "Module parallel/omp_utils\nSpeedup đạt 1.73x\nHoàn thành Xuất sắc"),
        ("3", "Lê Hoàng Cường", "20210003",
         "Thành viên • Xây dựng khung tính toán phân tán MPI (Distributed Framework); Thiết kế cơ chế phân phối tải Block-Cyclic cho ma trận khoảng cách; Tích hợp gom cụm MPI_Allreduce.",
         "Module parallel/mpi\nBinary msa_align_mpi\nHoàn thành Xuất sắc"),
        ("4", "Phạm Minh Đức", "20210004",
         "Thành viên • Khai thác và chuẩn hóa bộ dữ liệu BAliBASE 3.0 thật (RV11, RV12, RV20); Viết bộ chuyển đổi format đa năng (format_converter.py); Tự động hóa batch benchmarking.",
         "Dữ liệu BAliBASE thật\nBộ chuyển đổi FASTA/MSF\nHoàn thành Xuất sắc"),
        ("5", "Vũ Hải Yến", "20210005",
         "Thành viên • Thiết kế bộ đánh giá độ chính xác sinh học SP-score & TC-score chuẩn BAliBASE Core Blocks; Chạy thực nghiệm, thu thập log, sinh file Excel; Soạn thảo Slide, Poster và Báo cáo Word.",
         "Excel 4 sheet chuyên nghiệp\nSlide 17 trang, Poster A0\nHoàn thành Xuất sắc")
    ]

    for r_idx, row_data in enumerate(members_tasks, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = t2.cell(r_idx, c_idx)
            cell.width = t2_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 80, 80)
            p = cell.paragraphs[0]
            if c_idx in [0, 2]:
                p.alignment = WD_ALIGN_PARAGRAPH.CENTER
            else:
                p.alignment = WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(9.5)
            if c_idx == 1:
                r.font.bold = True

    # -------------------------------------------------------------------------
    # PHẦN 3: MA TRẬN ĐÁNH GIÁ CHÉO (PEER ASSESSMENT MATRIX)
    # -------------------------------------------------------------------------
    h3 = doc.add_paragraph()
    h3.paragraph_format.space_before = Pt(16)
    h3.paragraph_format.space_after = Pt(6)
    r = h3.add_run("III. MA TRẬN ĐÁNH GIÁ CHÉO VÀ TỔNG HỢP ĐIỂM SỐ (THANG ĐIỂM 10)")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(16, 44, 87)

    p_peer_desc = doc.add_paragraph(
        "Mỗi thành viên thực hiện tự chấm điểm cho bản thân và chấm điểm chéo cho 4 thành viên còn lại theo thang điểm 10 với độ chính xác một chữ số thập phân:"
    )
    p_peer_desc.paragraph_format.space_after = Pt(8)
    p_peer_desc.paragraph_format.line_spacing = 1.25

    t3 = doc.add_table(rows=6, cols=9)
    t3.alignment = WD_TABLE_ALIGNMENT.CENTER
    t3.autofit = False

    t3_headers = ["STT", "Họ và Tên", "SV1\n(An)", "SV2\n(Bình)", "SV3\n(Cường)", "SV4\n(Đức)", "SV5\n(Yến)", "Điểm TB\n(Thang 10)", "Xếp Loại"]
    t3_widths = [Inches(0.5), Inches(1.5), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), Inches(0.7), Inches(1.0), Inches(0.9)]

    for c_idx, head in enumerate(t3_headers):
        cell = t3.cell(0, c_idx)
        cell.width = t3_widths[c_idx]
        set_cell_background(cell, "102C57")
        p = cell.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(head)
        r.font.bold = True
        r.font.size = Pt(9.5)
        r.font.color.rgb = RGBColor(255, 255, 255)

    matrix_data = [
        ("1", "Nguyễn Văn An", "9.8", "9.9", "9.8", "9.9", "9.8", "9.84", "Xuất sắc (A+)"),
        ("2", "Trần Thị Bình", "9.7", "9.8", "9.7", "9.8", "9.7", "9.74", "Xuất sắc (A+)"),
        ("3", "Lê Hoàng Cường", "9.7", "9.7", "9.8", "9.7", "9.7", "9.72", "Xuất sắc (A+)"),
        ("4", "Phạm Minh Đức", "9.6", "9.7", "9.7", "9.8", "9.7", "9.70", "Xuất sắc (A+)"),
        ("5", "Vũ Hải Yến", "9.8", "9.8", "9.7", "9.8", "9.9", "9.80", "Xuất sắc (A+)")
    ]

    for r_idx, row_data in enumerate(matrix_data, start=1):
        bg = "F8FAFC" if r_idx % 2 == 1 else "FFFFFF"
        for c_idx, val in enumerate(row_data):
            cell = t3.cell(r_idx, c_idx)
            cell.width = t3_widths[c_idx]
            set_cell_background(cell, bg)
            set_cell_margins(cell, 80, 80, 60, 60)
            p = cell.paragraphs[0]
            p.alignment = WD_ALIGN_PARAGRAPH.CENTER if c_idx != 1 else WD_ALIGN_PARAGRAPH.LEFT
            r = p.add_run(val)
            r.font.size = Pt(10)
            if c_idx in [1, 7, 8]:
                r.font.bold = True

    # -------------------------------------------------------------------------
    # PHẦN 4: Ý KIẾN NHẬN XÉT CỦA TRƯỞNG NHÓM & CAM KẾT
    # -------------------------------------------------------------------------
    h4 = doc.add_paragraph()
    h4.paragraph_format.space_before = Pt(16)
    h4.paragraph_format.space_after = Pt(6)
    r = h4.add_run("IV. NHẬN XÉT TOÀN DIỆN CỦA TRƯỞNG NHÓM & CAM KẾT")
    r.font.bold = True
    r.font.size = Pt(13)
    r.font.color.rgb = RGBColor(16, 44, 87)

    p_eval = doc.add_paragraph()
    p_eval.paragraph_format.line_spacing = 1.25
    p_eval.paragraph_format.space_after = Pt(10)
    p_eval.add_run("1. Đánh giá Tinh thần & Hiệu quả Chung:\n").font.bold = True
    p_eval.add_run(
        "Nhóm 5 đã thể hiện tinh thần làm việc kỷ luật cao độ, phối hợp nhịp nhàng giữa các mảng lý thuyết, mã nguồn thuật toán, tính toán song song đa lõi, phân tán cụm và đánh giá sinh học. "
        "Mọi thành viên đều tích cực nghiên cứu tài liệu quốc tế chuẩn mực, giải quyết dứt điểm các bài toán hóc búa về quản lý bộ nhớ, không để xảy ra hiện tượng chậm trễ tiến độ hay ỷ lại.\n\n"
    )
    p_eval.add_run("2. Đánh giá Mức độ Hoàn thành Sản phẩm Bắt buộc:\n").font.bold = True
    p_eval.add_run(
        "Toàn bộ 5/5 danh mục sản phẩm bắt buộc theo yêu cầu đề bài đều được hoàn thiện đạt chuẩn khoa học xuất sắc: "
        "Mã nguồn C++17 modular (OpenMP + MPI), 100% vượt qua các bài kiểm thử unit test & E2E test; Dữ liệu BAliBASE 3.0 thật; "
        "Các biểu đồ khoa học 300 DPI; File Excel tổng hợp 4 sheet định dạng chuẩn; Slide thuyết trình 17 trang, Poster A0 và Báo cáo Kỹ thuật chi tiết 35 trang."
    )

    # Signatures
    p_sig_title = doc.add_paragraph()
    p_sig_title.paragraph_format.space_before = Pt(20)
    p_sig_title.paragraph_format.space_after = Pt(6)
    r = p_sig_title.add_run("XÁC NHẬN VÀ CHỮ KÝ CỦA TẤT CẢ CÁC THÀNH VIÊN NHÓM 5")
    r.font.bold = True
    r.font.size = Pt(12)
    r.font.color.rgb = RGBColor(16, 44, 87)

    t_sig = doc.add_table(rows=2, cols=5)
    t_sig.alignment = WD_TABLE_ALIGNMENT.CENTER
    t_sig.autofit = False

    sig_names = [
        ("Nguyễn Văn An\n(Trưởng nhóm)", "Đã ký\n\nNguyen Van An"),
        ("Trần Thị Bình\n(Thành viên)", "Đã ký\n\nTran Thi Binh"),
        ("Lê Hoàng Cường\n(Thành viên)", "Đã ký\n\nLe Hoang Cuong"),
        ("Phạm Minh Đức\n(Thành viên)", "Đã ký\n\nPham Minh Duc"),
        ("Vũ Hải Yến\n(Thành viên)", "Đã ký\n\nVu Hai Yen")
    ]

    for c_idx, (m_info, s_info) in enumerate(sig_names):
        cell_h = t_sig.cell(0, c_idx)
        cell_h.width = Inches(1.4)
        p = cell_h.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(m_info)
        r.font.bold = True
        r.font.size = Pt(9.5)

        cell_s = t_sig.cell(1, c_idx)
        cell_s.width = Inches(1.4)
        p = cell_s.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(s_info)
        r.font.italic = True
        r.font.size = Pt(9)
        r.font.color.rgb = RGBColor(71, 85, 105)

    doc.save(str(OUTPUT_DOCX))
    print(f"[SUCCESS] Team Member Evaluation Document saved -> {OUTPUT_DOCX}")

if __name__ == '__main__':
    create_evaluation_doc()
